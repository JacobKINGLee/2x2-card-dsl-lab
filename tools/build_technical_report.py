from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\Generative UI")
OUTPUT = ROOT / "2x2卡片语义DSL与约束布局引擎_技术报告.docx"

# standard_business_brief preset + named CJK font override
FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
FONT_MONO = "Consolas"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
INK = "20242B"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
BORDER = "D0D5DD"
WHITE = "FFFFFF"
GREEN = "157A58"
RED = "9B1C1C"
GOLD = "7A5A00"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name=FONT_LATIN, east_asia=FONT_CJK, size=None,
                 color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = edges.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_paragraph(paragraph, fill=CALLOUT, border_color=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border_color:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        for edge in ("top", "left", "bottom", "right"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "8")
            node.set(qn("w:space"), "4")
            node.set(qn("w:color"), border_color)
            p_bdr.append(node)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def set_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=11, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


_current_decimal_num_id = None


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id:08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(lvl)
    # OOXML requires every abstractNum to precede the concrete num instances.
    # Inserting at the end makes Word repair the numbering part and can corrupt
    # built-in bullet styles.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def add_number(doc, text, restart=False):
    global _current_decimal_num_id
    if restart or _current_decimal_num_id is None:
        _current_decimal_num_id = create_decimal_numbering(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(_current_decimal_num_id))
    num_pr.extend([ilvl, num_id])
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_callout(doc, label, text, color=BLUE, fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(
        cell,
        left={"val": "single", "sz": "24", "color": color},
        top={"val": "single", "sz": "4", "color": BORDER},
        bottom={"val": "single", "sz": "4", "color": BORDER},
        right={"val": "single", "sz": "4", "color": BORDER},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, color=color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code_block(doc, code, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(caption)
        set_run_font(r, size=9, color=MUTED, bold=True)
        set_keep_with_next(p)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    shade_paragraph(p, "111827", "374151")
    lines = code.strip("\n").splitlines()
    for index, line in enumerate(lines):
        r = p.add_run(line)
        set_run_font(r, name=FONT_MONO, east_asia=FONT_CJK, size=8.2, color="F8FAFC")
        if index < len(lines) - 1:
            r.add_break()
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0 and len(headers) > 1:
                r = p.add_run(str(value))
                set_run_font(r, size=font_size, color=DARK_BLUE, bold=True)
            else:
                r = p.add_run(str(value))
                set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return table


def add_flow_step(doc, number, name, input_output):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [700, 2500, 6160])
    table.style = "Table Grid"
    cells = table.rows[0].cells
    set_cell_shading(cells[0], BLUE)
    set_cell_shading(cells[1], BLUE_GRAY)
    set_cell_shading(cells[2], WHITE)
    for idx, text in enumerate((number, name, input_output)):
        p = cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_run_font(
            r,
            size=9.5 if idx != 1 else 10,
            color=WHITE if idx == 0 else NAVY if idx == 1 else INK,
            bold=idx < 2,
        )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("↓")
    set_run_font(r, size=9, color=MUTED, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def add_result_placeholder(doc, title, prompt, min_height_inches):
    table = doc.add_table(rows=2, cols=1)
    set_table_geometry(table, [9360])
    header = table.cell(0, 0)
    set_cell_shading(header, BLUE_GRAY)
    set_cell_border(header,
                    top={"val": "single", "sz": "6", "color": BORDER},
                    left={"val": "single", "sz": "6", "color": BORDER},
                    right={"val": "single", "sz": "6", "color": BORDER},
                    bottom={"val": "single", "sz": "4", "color": BORDER})
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run(title)
    set_run_font(hr, size=10.5, color=NAVY, bold=True)

    body = table.cell(1, 0)
    set_cell_shading(body, "FBFCFD")
    set_cell_border(body,
                    top={"val": "single", "sz": "4", "color": BORDER},
                    left={"val": "single", "sz": "6", "color": BORDER},
                    right={"val": "single", "sz": "6", "color": BORDER},
                    bottom={"val": "single", "sz": "6", "color": BORDER})
    row = table.rows[1]
    row.height = Inches(min_height_inches)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    bp = body.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bp.paragraph_format.space_before = Pt(8)
    br = bp.add_run(prompt)
    set_run_font(br, size=9, color="98A2B3", italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc, force=False):
    """Add a hard page break only for deliberately isolated sections.

    For ordinary chapter transitions, Word's keep-with-next heading behavior is
    enough and avoids leaving a mostly empty page when a table happens to split.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    if force:
        p.add_run().add_break(WD_BREAK.PAGE)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_styles(doc)

    # Quiet running furniture for a multi-page technical report.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("2×2 卡片约束布局引擎｜技术报告")
    set_run_font(hr, size=8.5, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])

    # Cover / memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("技术报告")
    set_run_font(r, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("面向 2×2 卡片的语义 DSL\n与约束布局引擎")
    set_run_font(r, size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("从无坐标元素描述到可渲染 Layout IR 的端到端原型")
    set_run_font(r, size=13, color=MUTED)

    metadata = [
        ("文档性质", "阶段性技术方案与实现说明"),
        ("提交对象", "导师评审"),
        ("项目范围", "160×160vp（2×2）卡片的语义 DSL、约束求解与渲染"),
        ("当前阶段", "候选布局枚举 + 硬约束校验 + 软约束评分 + 250 组压力测试"),
        ("版本日期", "2026-08-27"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, color=NAVY, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)

    add_heading(doc, "摘要", 1)
    add_body(doc, "本项目验证一种不依赖固定卡片模板的生成式 UI 路径：上游仅输出少量语义元素，布局引擎依据 UX 规范生成候选矩形，对候选执行硬约束校验与软约束评分，最后输出包含明确坐标的 LayoutResult，并由 Renderer 完成 160×160vp 卡片渲染。当前实现是一个面向 2×2 卡片场景的专用约束布局引擎，能够处理文本、应用图标、图标按钮、胶囊按钮、指标与图片等元素，并支持文本压缩、可选元素舍弃和随机组合压力测试。")
    add_callout(doc, "核心结论", "当前 Demo 已经证明“语义 DSL → 自动坐标求解 → 可验证渲染”链路可行；它不是固定模板拼装，也不是完全通用的数学求解器，而是介于二者之间的领域专用搜索与评分系统。")

    page_break(doc, force=True)

    add_heading(doc, "1. 研究背景与问题定义", 1)
    add_heading(doc, "1.1 背景", 2)
    add_body(doc, "传统卡片 UI 通常采用预制模板：为某类卡片提前写死标题区、内容区和按钮区的位置，再将业务数据填入模板。该方式实现简单、视觉稳定，但当元素组合变化时，需要持续增加模板数量，并且难以直接验证“任意少量元素是否都能自动获得满足 UX 的坐标”。")
    add_body(doc, "导师提出的研究方向是：在卡片尺寸固定、元素数量较少、变化空间有限的前提下，是否可以取消模板选择，让模型或业务层只输出元素语义，再由布局系统给出满足 UX 条件的坐标结果。")

    add_heading(doc, "1.2 研究问题", 2)
    add_bullet(doc, "如何约束上游 DSL，使其描述内容语义而不能直接写入坐标和样式？")
    add_bullet(doc, "如何根据元素类型、文本长度、按钮位置和并行指标数量生成有限且可解释的候选布局？")
    add_bullet(doc, "如何区分不能违反的 UX 规则与可用于比较候选优劣的偏好规则？")
    add_bullet(doc, "如何输出稳定的 Layout IR，使同一求解结果可以被 Web、端侧或其他 Renderer 消费？")

    add_heading(doc, "1.3 本阶段目标", 2)
    add_number(doc, "建立仅包含语义元素的 Element DSL。", restart=True)
    add_number(doc, "实现 160×160vp 卡片的候选布局生成与坐标分配。")
    add_number(doc, "实现安全区、元素重叠、区域间距、固定控件尺寸和必选元素放置等硬约束。")
    add_number(doc, "实现基于 priority、信息保留率和布局偏好的软约束评分。")
    add_number(doc, "通过 CoordinateCard 展示最终结果，并通过固定随机种子生成 250 组组合进行自动化评估。")

    add_heading(doc, "1.4 UX 规范来源与实现假设", 2)
    add_table(doc,
              ["UX条目", "当前实现", "性质"],
              [
                  ("卡片尺寸", "160×160vp", "实现假设"),
                  ("安全边距", "四周12vp", "原始规范"),
                  ("卡片圆角", "20vp", "原始规范"),
                  ("应用图标", "20×20vp，右上角", "原始规范"),
                  ("标题与图标", "标题宽112vp，为图标预留4vp", "原始规范"),
                  ("按钮", "36vp高；图标按钮36×36vp", "规范+补充假设"),
                  ("区域间距", "标题/内容/操作区主要使用8vp", "原始规范"),
                  ("指标单元最小值", "40×38vp", "实现假设"),
                  ("图片最小高度", "44vp", "实现假设"),
              ],
              [1800, 4800, 2760])
    add_callout(doc, "边界说明", "未在第一版 UX 文档中明确规定的尺寸和取舍规则，以可解释、可替换的引擎常量实现；后续可在与导师确认后转为正式设计 Token。", color=GOLD, fill="FFF8E8")

    page_break(doc)

    add_heading(doc, "2. 总体方案与端到端链路", 1)
    add_body(doc, "系统被划分为语义层、求解层和渲染层。语义层保证输入干净；求解层负责从多个候选中选择合法且信息保留较好的坐标结果；渲染层只消费结果，不参与布局决策。")

    pipeline = [
        ("01", "JSON字符串", "编辑器中的原始文本；可能存在语法或字段错误"),
        ("02", "parseElementDSL", "JSON解析 + 运行时语义校验"),
        ("03", "Element DSL", "无坐标、无样式的受控语义元素集合"),
        ("04", "solveLayout", "建立展示方案、按钮锚点和指标列数搜索空间"),
        ("05", "候选布局", "attemptLayout 为每个组合生成 x / y / width / height"),
        ("06", "硬约束检查", "安全区、重叠、间距、控件尺寸、必选元素"),
        ("07", "软约束评分", "信息保留率、priority、默认锚点、指标形状"),
        ("08", "LayoutResult", "最佳候选、坐标节点、评分、决策与失败原因"),
        ("09", "CoordinateCard", "将坐标映射为 CSS 绝对定位并渲染真实元素"),
    ]
    for idx, (number, name, description) in enumerate(pipeline):
        add_flow_step(doc, number, name, description)
        if idx == len(pipeline) - 1:
            # Remove the final decorative arrow paragraph.
            last_p = doc.paragraphs[-1]
            last_p._element.getparent().remove(last_p._element)

    add_heading(doc, "2.1 功能模块划分", 2)
    add_table(doc,
              ["模块", "主要职责", "当前代码位置"],
              [
                  ("DSL编辑与预设", "维护 JSON 字符串、切换示例、格式化输入", "app/page.tsx"),
                  ("DSL解析器", "语法解析、字段校验、数量约束、禁止坐标/样式", "app/layout-engine.ts / parseElementDSL"),
                  ("文本测量", "Canvas真实测量或字符宽度估算，计算行数和高度", "app/layout-engine.ts / measureTextBlock"),
                  ("展示策略生成", "生成 full、compact、hidden 组合", "app/layout-engine.ts / buildPresentationPlans"),
                  ("候选坐标生成", "为按钮、图标、文本、图片和指标分配矩形", "app/layout-engine.ts / attemptLayout"),
                  ("约束与评分", "过滤违规候选、计算信息保留效用与偏好得分", "app/layout-engine.ts"),
                  ("结果选择", "优先选择得分最高的 solved 候选", "app/layout-engine.ts / solveLayout"),
                  ("坐标渲染器", "将 LayoutNode 坐标转换为 CSS left/top/width/height", "app/page.tsx / CoordinateCard"),
                  ("自动评估", "固定种子生成250组文本、指标和图片组合", "app/layout-benchmark.ts"),
              ],
              [1800, 4140, 3420], font_size=8.8)

    page_break(doc)

    add_heading(doc, "3. 语义层：从 JSON 字符串到 Element DSL", 1)
    add_heading(doc, "3.1 输入数据模型", 2)
    add_code_block(doc, '''export type ElementDSL = {
  version: "2.0" | "3.0";
  type: "adaptive-card";
  elements: CardElement[];
};

type CardElement =
  | TextElement | AppIconElement
  | IconButtonElement | CapsuleButtonElement
  | MetricElement | ImageElement;''', "代码摘录：Element DSL 的核心类型")
    add_body(doc, "Element DSL 只承担“内容契约”的角色。元素可以声明 id、type、文本、事件、priority、optional 和有限的语义偏好，但不能声明 x、y、width、height 或 style。这样可以避免上游绕过 UX 规则，也使同一份语义输入可以在不同 Renderer 中复用。")

    add_heading(doc, "3.2 parseElementDSL 的职责", 2)
    add_table(doc,
              ["校验层级", "典型规则", "失败结果"],
              [
                  ("JSON语法", "JSON.parse 能否成功", "data=null，返回语法错误"),
                  ("根节点", "必须是对象；version/type合法", "记录根级错误"),
                  ("元素数组", "1–8个元素，每个id唯一", "记录数量或重复id错误"),
                  ("禁用字段", "禁止style/layout与元素坐标尺寸", "拒绝进入布局求解"),
                  ("类型字段", "不同type要求对应的必填字段", "记录具体字段路径"),
                  ("全局基数", "最多一个操作、一个图标、一张图片、四个指标", "记录组合错误"),
                  ("互斥关系", "图片与指标当前不能同时出现", "记录场景不支持"),
              ],
              [1800, 4720, 2840])

    add_code_block(doc, '''if ("x" in element || "y" in element ||
    "width" in element || "height" in element ||
    "style" in element) {
  errors.push("不能直接声明坐标、尺寸或样式");
}''', "关键控制：禁止输入端直接决定布局")

    add_heading(doc, "3.3 priority 与 optional", 2)
    add_body(doc, "priority 是 0–100 的语义重要度。未提供时，引擎根据元素类型补充默认值：操作按钮100、标题90、正文80、指标75、应用图标65、说明文字50、图片及其他元素40。")
    add_bullet(doc, "priority 不直接决定 x/y 坐标。")
    add_bullet(doc, "较低 priority 的可降级文本更早进入 compact。")
    add_bullet(doc, "较低 priority 且 optional=true 的元素更早进入 hidden。")
    add_bullet(doc, "最终评分按 priority 计算被完整保留、压缩、截断或舍弃的信息价值。")
    add_callout(doc, "解释性", "同一布局为什么压缩某段文本、舍弃某个说明元素，可以通过 priority、compressedIds、droppedIds 和 decisions 追溯。")

    page_break(doc)

    add_heading(doc, "4. 求解层：候选空间的生成", 1)
    add_heading(doc, "4.1 求解器入口", 2)
    add_code_block(doc, '''export function solveLayout(
  dsl: ElementDSL,
  measurer: TextMeasurer = estimatedTextMeasurer,
): LayoutResult''', "代码摘录：solveLayout")
    add_body(doc, "solveLayout 不直接返回第一种可用布局，而是先建立有限搜索空间。当前候选数量由三个维度相乘得到：展示方案数量 × 图标按钮锚点数量 × 可见指标列数数量。")

    add_heading(doc, "4.2 展示方案：full / compact / hidden", 2)
    add_table(doc,
              ["模式", "含义", "触发方式", "评分保留系数"],
              [
                  ("full", "完整显示主文本与 supporting", "基础方案", "1.00"),
                  ("compact", "正文最多一行，不显示 supporting", "按低 priority 逐步压缩", "0.72"),
                  ("truncated", "节点存在但文本超出最大行数", "测量结果超过 maxLines", "0.86"),
                  ("hidden", "元素不进入坐标节点", "仅 optional 元素可被舍弃", "0.00"),
              ],
              [1500, 3020, 3260, 1580])
    add_body(doc, "buildPresentationPlans 首先创建所有元素均为 full 的基础方案，再将可降级文本按优先级从低到高分别或累计压缩，最后对 optional 元素分别或累计隐藏。方案使用 planKey 去重，并限制在最多64种，以避免组合爆炸。")

    add_heading(doc, "4.3 图标按钮锚点", 2)
    add_table(doc,
              ["placement输入", "候选锚点", "说明"],
              [
                  ("bottom-start", "仅左下", "显式语义偏好，软评分获得明确指定奖励"),
                  ("bottom-end", "仅右下", "显式语义偏好"),
                  ("auto", "右下 + 左下", "同时求解；两者都合法时默认略偏右下"),
                  ("无图标按钮", "占位值bottom-end", "不产生实际操作节点"),
              ],
              [1900, 2300, 5160])

    add_heading(doc, "4.4 指标列数", 2)
    add_body(doc, "当存在 N 个可见 metric 时，求解器枚举1至N列。以3个指标为例，会比较1列×3行、2列×2行、3列×1行。每种网格都必须满足单元格至少40×38vp，并在软评分中倾向宽高比例更均衡的结果。")

    add_code_block(doc, '''presentationPlans.forEach((plan) => {
  placements.forEach((placement) => {
    columnChoices.forEach((columns) => {
      candidates.push(
        attemptLayout(dsl, placement, plan, columns, measurer)
      );
    });
  });
});''', "代码摘录：候选布局三层枚举")

    page_break(doc)

    add_heading(doc, "5. 求解层：单个候选的坐标生成", 1)
    add_heading(doc, "5.1 坐标系与基础常量", 2)
    add_table(doc,
              ["常量", "值", "作用"],
              [
                  ("CARD_SIZE", "160", "卡片宽高"),
                  ("SAFE_MARGIN", "12", "四周安全边距"),
                  ("INNER_SIZE", "136", "安全区宽度：160-12×2"),
                  ("REGION_GAP", "8", "标题、内容和操作区主要间距"),
              ],
              [2100, 1500, 5760])

    add_heading(doc, "5.2 放置顺序", 2)
    add_number(doc, "优先放置底部操作按钮，并据此确定内容区底边 contentBottom。", restart=True)
    add_number(doc, "放置右上角应用图标。")
    add_number(doc, "按 DSL 中的文本语义顺序，从安全区顶部向下排列文本。")
    add_number(doc, "在文本与操作区之间的剩余矩形中放置图片或指标网格。")
    add_number(doc, "检查必选元素、安全区和两两重叠。")

    add_heading(doc, "5.3 固定控件坐标", 2)
    add_table(doc,
              ["元素", "坐标规则", "典型结果"],
              [
                  ("应用图标", "x=160-12-20，y=12，20×20", "(128, 12, 20, 20)"),
                  ("左下图标按钮", "x=12，y=160-12-36，36×36", "(12, 112, 36, 36)"),
                  ("右下图标按钮", "x=160-12-36，y=112，36×36", "(112, 112, 36, 36)"),
                  ("胶囊按钮", "x=12，y=112，w=136，h=36", "(12, 112, 136, 36)"),
              ],
              [2100, 4440, 2820])

    add_heading(doc, "5.4 文本测量与垂直流", 2)
    add_body(doc, "页面等待 document.fonts.ready 后创建 Canvas TextMeasurer，通过 context.measureText 获取文本宽度；服务端或字体未就绪时使用字符宽度估算器。measureTextBlock 将测量宽度除以可用宽度得到自然行数，再结合 role、maxLines 和 presentation 计算实际行数、高度及是否截断。")
    add_bullet(doc, "与右上角图标同一行时，文本宽度为112vp；否则使用完整安全区宽136vp。")
    add_bullet(doc, "标题或相邻标题与文本之间使用8vp间距；其他文本之间当前使用4vp。")
    add_bullet(doc, "body 在 full 模式下最多3行；title和caption最多1行；compact统一最多1行。")
    add_bullet(doc, "supporting 只在 full 模式显示，并增加16vp高度。")

    add_heading(doc, "5.5 图片与指标的剩余空间", 2)
    add_body(doc, "图片填充约束求解后的最大剩余矩形，剩余高度不足44vp时该候选失败。指标根据当前列数均分安全区宽度和剩余高度，列间、行间均使用8vp；单元格小于40×38vp时记录硬约束违反。")

    page_break(doc)

    page_break(doc, force=True)
    add_heading(doc, "6. 硬约束检查与软约束评分", 1)
    add_heading(doc, "6.1 硬约束：决定候选是否可用", 2)
    add_table(doc,
              ["检查项", "判定方式", "失败后果"],
              [
                  ("12vp安全区域", "所有节点完全位于[12,148]范围", "status=unsatisfied"),
                  ("元素不重叠", "所有LayoutNode矩形两两相交检测", "记录具体id对"),
                  ("内容与操作区≥8vp", "contentBottom=按钮顶部-8", "超高文本无法获得坐标"),
                  ("控件固定尺寸", "图标20×20；操作按钮高度36", "固定控件检查失败"),
                  ("必选元素全部布局", "optional!=true的元素必须出现在nodes", "记录未放置元素"),
                  ("图片/指标最小矩形", "图片≥44高；指标≥40×38", "候选产生violation"),
              ],
              [2200, 4580, 2580])
    add_body(doc, "硬约束是准入门槛。软评分不能将存在违规的候选“补救”为合法候选。只有 violations 为空，候选的 status 才是 solved。")

    add_heading(doc, "6.2 软约束：在合法候选中比较优劣", 2)
    add_code_block(doc, '''utility = retainedPriority / totalPriority
score = round(utility * 82)
score += noViolation ? 10 : 0
score += actionPreferenceBonus
score += metricAspectBonus
score -= violations.length * 22
score = clamp(score, 0, 100)''', "当前评分公式的概念化表示")
    add_table(doc,
              ["评分组成", "作用"],
              [
                  ("信息保留效用", "高 priority 元素完整保留时得分更高；压缩、截断和舍弃会产生不同损失"),
                  ("合法奖励", "无违规候选增加10分"),
                  ("按钮偏好", "显式placement或auto右下获得5分；auto左下获得3分"),
                  ("指标形状", "指标单元格越接近均衡比例，奖励越高"),
                  ("违规惩罚", "每项violation扣22分，主要用于无合法结果时比较“最少违规”候选"),
              ],
              [2200, 7160])

    add_heading(doc, "6.3 最佳候选选择", 2)
    add_number(doc, "先过滤出 status=solved 的候选。", restart=True)
    add_number(doc, "若存在合法候选，按 score 从高到低选择第一名。")
    add_number(doc, "若没有合法候选，先按 violations 数量升序，再按 score 降序返回一个最接近可行的诊断候选。")
    add_callout(doc, "重要区分", "诊断候选仍保持 unsatisfied。主预览不会把它当成正式布局；压力测试浏览器仅用灰色 REJECTED CANDIDATE 形式展示，帮助定位失败原因。", color=RED, fill="FDF2F2")

    page_break(doc)

    add_heading(doc, "7. 输出层：LayoutResult、Layout IR 与 CoordinateCard", 1)
    add_heading(doc, "7.1 LayoutResult 数据结构", 2)
    add_code_block(doc, '''type LayoutResult = {
  status: "solved" | "unsatisfied";
  card: { width: number; height: number; safeMargin: number };
  nodes: LayoutNode[];
  violations: string[];
  decisions: string[];
  checks: ConstraintCheck[];
  candidateCount: number;
  score: number;
  compressedIds: string[];
  droppedIds: string[];
  measurement: "canvas" | "estimate";
};''', "代码摘录：求解器公开结果")
    add_body(doc, "LayoutResult 同时服务于渲染和解释。nodes 是 Renderer 的必要输入；score、checks、violations、compressedIds、droppedIds 和 decisions 则用于调试、评审和实验统计。")

    add_heading(doc, "7.2 LayoutNode 与 Layout IR", 2)
    add_table(doc,
              ["字段", "作用"],
              [
                  ("id", "与输入Element DSL中的元素建立稳定对应"),
                  ("x / y", "元素左上角相对于卡片左上角的坐标"),
                  ("width / height", "元素的求解矩形"),
                  ("element", "原始语义元素，供Renderer选择组件并填充内容"),
                  ("reason", "当前坐标的生成原因"),
                  ("presentation", "full或compact"),
                  ("truncated / lineCount", "控制文本截断和行数"),
              ],
              [2200, 7160])
    add_body(doc, "页面还会从 LayoutResult 提取轻量 Layout IR，只保留卡片尺寸、评分、测量方式、压缩/舍弃元素以及节点坐标。该结构可以作为跨端渲染协议的雏形。")

    add_heading(doc, "7.3 CoordinateCard 的职责边界", 2)
    add_code_block(doc, '''const style = {
  left: node.x,
  top: node.y,
  width: node.width,
  height: node.height,
};

<div className="layout-node" style={style}>
  {renderElement(node, onAction)}
</div>''', "代码摘录：坐标到CSS绝对定位的映射")
    add_bullet(doc, "CoordinateCard 不重新计算坐标，只消费 LayoutResult。")
    add_bullet(doc, "卡片容器为 position:relative、160×160px；节点为 position:absolute。")
    add_bullet(doc, "当前 Demo 采用 1vp≈1px 的模拟，再用 transform:scale(1.72) 放大预览。")
    add_bullet(doc, "renderElement 根据 type 渲染文本、图标、按钮、指标或图片，并通过 event 触发交互。")

    page_break(doc)

    add_heading(doc, "8. 端到端示例：待办卡片", 1)
    add_heading(doc, "8.1 输入 Element DSL", 2)
    add_code_block(doc, '''{
  "version": "3.0",
  "type": "adaptive-card",
  "elements": [
    {"id":"title","type":"text","role":"title",
     "text":"待办事项","priority":100},
    {"id":"app","type":"appIcon","symbol":"办",
     "label":"待办应用","priority":90},
    {"id":"date","type":"text","role":"caption",
     "text":"今天","priority":70},
    {"id":"task","type":"text","role":"body",
     "text":"完成方案评审","supporting":"14:30 · 会议室 A",
     "priority":95},
    {"id":"complete","type":"iconButton","icon":"✓",
     "label":"标记为完成","event":"completeTodo",
     "placement":"auto","priority":100}
  ]
}''')

    add_heading(doc, "8.2 候选生成", 2)
    add_body(doc, "该示例包含一个可降级 body 文本、一个 placement=auto 的图标按钮且无指标，因此主要形成2种展示方案（full/compact）×2个按钮锚点（右下/左下）×1种列数选择，共约4个候选。")

    add_heading(doc, "8.3 典型坐标结果", 2)
    add_table(doc,
              ["元素", "X", "Y", "W", "H", "生成原因"],
              [
                  ("title", "12", "12", "112", "20", "右上图标占位，标题缩至112vp"),
                  ("app", "128", "12", "20", "20", "应用图标固定右上"),
                  ("date", "12", "40", "136", "16", "标题后8vp"),
                  ("task", "12", "60", "136", "36", "语义顺序排列并显示supporting"),
                  ("complete", "112", "112", "36", "36", "auto比较后选择右下锚点"),
              ],
              [1600, 700, 700, 700, 700, 4960], font_size=8.8)

    add_heading(doc, "8.4 为什么选择右下", 2)
    add_body(doc, "左下和右下候选均能通过安全区、间距和重叠检查时，auto 模式对右下候选增加5分、左下候选增加3分，因此右下成为最佳布局。如果 DSL 显式指定 bottom-start，搜索空间只保留左下候选，按钮坐标将变为(12,112,36,36)。")
    add_callout(doc, "示例结论", "按钮位于左下并不需要创建新模板，只需改变语义偏好；引擎仍负责生成坐标并执行同一组约束。")

    page_break(doc)

    add_heading(doc, "9. 自动化评估与可观测性", 1)
    add_heading(doc, "9.1 250组固定种子压力测试", 2)
    add_body(doc, "layout-benchmark.ts 使用固定随机种子20260826生成250组样本，按索引循环覆盖文本组合、指标组合和图片组合。固定种子保证每次刷新得到相同样本，便于回归比较。")
    add_table(doc,
              ["场景", "随机变化"],
              [
                  ("文本组合", "标题、正文长度、supporting、可选caption、可选低优先级正文、可选应用图标与操作按钮"),
                  ("指标组合", "1–4个metric、可选标题、可选图标按钮、不同标签和值"),
                  ("图片组合", "标题、图片、可选caption、可选胶囊按钮"),
              ],
              [1900, 7460])

    add_heading(doc, "9.2 评估指标", 2)
    add_table(doc,
              ["指标", "定义"],
              [
                  ("Solve Rate", "250组中通过全部硬约束的比例"),
                  ("Average Score", "所有返回结果的平均软评分"),
                  ("Average Candidates", "每组输入平均生成的候选数量"),
                  ("Compressed Layouts", "使用compact或文本截断的样本数量"),
                  ("Dropped Layouts", "舍弃optional元素的样本数量"),
                  ("Violation Rate", "未找到合法候选的样本比例"),
                  ("Top Violations", "出现频率最高的三类失败原因"),
              ],
              [2200, 7160])

    add_heading(doc, "9.3 Demo中的可观测信息", 2)
    add_bullet(doc, "Layout Score、候选数量、硬约束通过数和文本测量方式。")
    add_bullet(doc, "compressedIds 与 droppedIds，用于观察优先级降级。")
    add_bullet(doc, "坐标表与完整 Layout IR JSON。")
    add_bullet(doc, "Solver Decisions，用自然语言记录本次坐标生成与降级原因。")
    add_bullet(doc, "Test Case Browser，可浏览第1–250组 Element DSL、结果、失败原因和 Renderer 画面。")

    add_heading(doc, "9.4 验收标准建议", 2)
    add_number(doc, "每次修改布局规则后，保持固定种子重新运行250组测试。", restart=True)
    add_number(doc, "比较求解成功率、平均分、压缩率、舍弃率及Top Violations是否退化。")
    add_number(doc, "人工抽查文本、指标、图片三类的 solved 与 rejected 样本。")
    add_number(doc, "将不符合视觉预期但通过硬约束的样本记录为软约束缺口，补充评分或候选生成策略。")

    page_break(doc)

    add_heading(doc, "10. 方案性质、局限与下一步", 1)
    add_heading(doc, "10.1 与固定模板方案的区别", 2)
    add_table(doc,
              ["比较维度", "固定模板", "当前约束布局引擎"],
              [
                  ("输入", "模板标识+业务数据", "无坐标语义元素"),
                  ("坐标来源", "模板代码预先写死", "运行时枚举和计算"),
                  ("元素组合", "依赖已存在模板", "在支持的元素类型和数量内自动组合"),
                  ("UX保证", "依靠模板开发者", "通过统一硬约束检查"),
                  ("信息取舍", "通常由模板逻辑固定", "按priority压缩或舍弃"),
                  ("解释性", "较弱", "输出候选数、评分、检查和决策"),
              ],
              [1900, 3480, 3980])

    add_heading(doc, "10.2 当前方案的准确定位", 2)
    add_callout(doc, "定位", "当前实现是领域专用的枚举式约束布局引擎：它不依赖卡片业务模板，但仍包含针对2×2卡片的元素放置策略和有限候选规则。")
    add_body(doc, "因此不能将其表述为“任意元素都能自由布局”的通用引擎。当前支持的自由度包括按钮左右锚点、文本展示层级、optional元素舍弃、指标列数和剩余空间分配；标题、图标、操作区等仍具有明确的领域规则。")

    add_heading(doc, "10.3 当前局限", 2)
    add_bullet(doc, "文本换行按总宽度估算自然行数，没有实现逐词/逐字符换行模拟。")
    add_bullet(doc, "候选坐标仍由程序化放置策略产生，尚未引入Cassowary、线性规划或通用约束满足算法。")
    add_bullet(doc, "固定尺寸与最小尺寸中包含尚待UX确认的实现假设。")
    add_bullet(doc, "评分权重为启发式参数，尚未通过设计师标注数据或用户研究校准。")
    add_bullet(doc, "当前仅针对160×160vp卡片，不覆盖其他卡片规格、RTL布局和无障碍字号放大。")

    add_heading(doc, "10.4 下一阶段建议", 2)
    add_number(doc, "将所有实现假设整理为可配置 Design Token，并与导师/设计规范逐项确认。", restart=True)
    add_number(doc, "增加真实文本排版测量、字体回退和不同语言测试。")
    add_number(doc, "扩展候选生成：文本与指标混排、更多锚点、可变元素数量和局部对齐策略。")
    add_number(doc, "建立标准样本集与人工视觉评分，用于调优软约束权重。")
    add_number(doc, "当程序化枚举难以继续扩展时，再评估Cassowary或整数/线性规划。")
    add_number(doc, "将 Layout IR 从Demo内部类型提升为有版本号的跨端协议，并增加Schema与回归测试。")

    add_heading(doc, "10.5 当前结论", 2)
    add_body(doc, "在2×2卡片的有限元素场景中，不使用业务模板、改由无坐标语义DSL驱动专用约束布局引擎是可行的。当前Demo已经覆盖从输入约束、候选生成、坐标分配、硬约束校验、软评分、结果解释到Renderer渲染的完整链路，并具备250组可重复压力测试。后续工作的重点不是继续增加业务模板，而是提升候选空间、测量精度、评分可信度和跨端协议稳定性。")

    page_break(doc, force=True)

    add_heading(doc, "11. DSL 结果展示区（提交前补充）", 1)
    add_body(doc, "本页预留给最终演示结果。建议选择一个具有代表性的样本，同时展示无坐标 Element DSL、求解后的 Layout IR、Renderer 画面和关键指标，从而直接证明“输入不含坐标，输出由引擎自动生成”。")

    add_result_placeholder(doc, "A. Element DSL 输入", "在此粘贴 DSL 代码截图或文本；需保留 elements、priority、optional、placement 等字段", 1.15)
    add_result_placeholder(doc, "B. Layout IR / 坐标结果", "在此粘贴坐标 JSON 或 Coordinate Table；突出 x、y、width、height 由引擎生成", 0.95)
    add_result_placeholder(doc, "C. Renderer 最终画面", "在此粘贴 Demo 渲染截图；建议开启坐标框并同时显示160×160vp标尺", 1.45)

    add_heading(doc, "11.1 结果记录", 2)
    add_table(doc,
              ["样本名称", "元素数", "候选数", "状态", "评分", "压缩/舍弃", "硬约束"],
              [("【填写】", "【填写】", "【填写】", "SOLVED / UNSAT", "【填写】", "【填写】", "通过数/总数")],
              [1800, 900, 900, 1350, 900, 1800, 1710], font_size=8.8)

    page_break(doc, force=True)

    add_heading(doc, "附录 A：实现文件与关键函数索引", 1)
    add_table(doc,
              ["文件/资料", "关键内容"],
              [
                  ("2x2_card_UX_design_spec.md", "第一版卡片UX规范：安全边距、标题、图标、按钮和内容区域"),
                  ("demo/app/layout-engine.ts:1–114", "常量、Element DSL、LayoutNode与LayoutResult类型"),
                  ("demo/app/layout-engine.ts:128–299", "parseElementDSL运行时校验"),
                  ("demo/app/layout-engine.ts:301–368", "文本测量、字体角色与行高计算"),
                  ("demo/app/layout-engine.ts:421–468", "buildPresentationPlans降级方案生成"),
                  ("demo/app/layout-engine.ts:471–803", "attemptLayout坐标生成、硬约束与评分"),
                  ("demo/app/layout-engine.ts:806–858", "solveLayout候选枚举与最佳结果选择"),
                  ("demo/app/page.tsx:264–361", "renderElement与CoordinateCard"),
                  ("demo/app/page.tsx:385–473", "页面状态、文本测量、解析/求解调用与Layout IR"),
                  ("demo/app/layout-benchmark.ts", "250组固定种子样本生成与统计"),
              ],
              [3600, 5760], font_size=8.5)

    add_heading(doc, "附录 B：关键术语", 1)
    add_table(doc,
              ["术语", "定义"],
              [
                  ("Element DSL", "无坐标、无样式的语义元素描述"),
                  ("Presentation Plan", "元素full、compact或hidden状态的组合"),
                  ("Candidate Layout", "某一展示方案、按钮锚点和指标列数组合生成的临时布局结果"),
                  ("Hard Constraint", "违反后候选不能作为正式布局的规则"),
                  ("Soft Constraint", "用于在多个合法候选之间比较偏好的评分项"),
                  ("LayoutNode", "一个元素的坐标、尺寸、展示状态及原始语义数据"),
                  ("LayoutResult", "求解器的完整输出，包括节点、状态、评分、检查与解释信息"),
                  ("Layout IR", "从LayoutResult中提取的轻量、可跨Renderer消费的坐标中间表示"),
                  ("CoordinateCard", "读取LayoutResult并将节点坐标映射到CSS定位的Web Renderer"),
              ],
              [2200, 7160], font_size=8.5)

    # Core properties and metadata.
    doc.core_properties.title = "面向2×2卡片的语义DSL与约束布局引擎"
    doc.core_properties.subject = "生成式UI约束布局引擎技术报告"
    doc.core_properties.author = "Generative UI Project"
    doc.core_properties.keywords = "Generative UI, DSL, Layout Engine, UX Constraints, Renderer"

    # Word keeps a required end paragraph after a final table. Make it tiny so
    # it remains on the appendix page instead of creating a blank final page.
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        tail = doc.paragraphs[-1]
        tail.paragraph_format.space_before = Pt(0)
        tail.paragraph_format.space_after = Pt(0)
        tail.paragraph_format.line_spacing = Pt(1)
        run = tail.add_run("\u200b")
        set_run_font(run, size=1, color=WHITE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
